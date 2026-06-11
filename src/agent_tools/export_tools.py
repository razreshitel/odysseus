"""export_document tool — render a markdown document to a real Office file
(.docx / .xlsx / .pptx).

Design: the model keeps writing markdown (its strength); this tool turns that
markdown into a binary Office file deterministically (the model never has to
format a binary blob, which it would butcher). Invoked ONLY when the user
explicitly names a non-markdown format — see the tool_index description; markdown
+ create_document stays the default for everything else.

Self-contained and additive: a new module, registered via TOOL_HANDLERS /
TOOL_TAGS, so a future rebase touches only those small registration lines. The
office libraries (python-docx, python-pptx, openpyxl) are pure-Python with no
system dependencies; if any is missing the tool returns a clear error instead of
crashing.
"""

from typing import Dict, List, Optional, Tuple
import json
import logging
import os
import re

logger = logging.getLogger(__name__)

# Misspelling-tolerant format aliases (the user explicitly listed exel, xsl …).
_FORMAT_ALIASES = {
    "docx": "docx", "doc": "docx", "word": "docx", "ms word": "docx", "msword": "docx",
    "xlsx": "xlsx", "xls": "xlsx", "xsl": "xlsx", "excel": "xlsx", "exel": "xlsx",
    "excell": "xlsx", "spreadsheet": "xlsx", "sheet": "xlsx",
    "pptx": "pptx", "ppt": "pptx", "powerpoint": "pptx", "power point": "pptx",
    "slides": "pptx", "slide": "pptx", "deck": "pptx", "presentation": "pptx",
}
_EXT = {"docx": ".docx", "xlsx": ".xlsx", "pptx": ".pptx"}


def _normalize_format(raw: str) -> Optional[str]:
    return _FORMAT_ALIASES.get((raw or "").strip().lower().lstrip("."))


def _slug(text: str, fallback: str = "export") -> str:
    s = re.sub(r"[^a-z0-9]+", "-", (text or "").strip().lower()).strip("-")
    return (s or fallback)[:60]


# ── Markdown helpers ───────────────────────────────────────────────────────
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _strip_inline(text: str) -> str:
    """Drop inline markdown emphasis/code markers for plain-text targets."""
    text = _BOLD_RE.sub(r"\1", text)
    text = text.replace("`", "").replace("**", "")
    return text.strip()


def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.count("|") >= 2


def _split_row(line: str) -> List[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def _parse_tables(md: str) -> List[List[List[str]]]:
    """Return every markdown table as rows-of-cells (separator row removed)."""
    tables, cur = [], []
    for line in md.splitlines():
        if _is_table_row(line):
            if _TABLE_SEP_RE.match(line):
                continue
            cur.append([_strip_inline(c) for c in _split_row(line)])
        else:
            if cur:
                tables.append(cur)
                cur = []
    if cur:
        tables.append(cur)
    return tables


# ── Renderers ──────────────────────────────────────────────────────────────
def _render_docx(md: str, title: str, path: str) -> None:
    from docx import Document as Docx
    from docx.shared import Pt

    doc = Docx()
    if title:
        doc.add_heading(title, level=0)

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        # Table block
        if _is_table_row(line):
            block = []
            while i < len(lines) and _is_table_row(lines[i]):
                if not _TABLE_SEP_RE.match(lines[i]):
                    block.append([_strip_inline(c) for c in _split_row(lines[i])])
                i += 1
            if block:
                cols = max(len(r) for r in block)
                t = doc.add_table(rows=0, cols=cols)
                t.style = "Light Grid Accent 1"
                for ri, row in enumerate(block):
                    cells = t.add_row().cells
                    for ci in range(cols):
                        cells[ci].text = row[ci] if ci < len(row) else ""
                        if ri == 0:
                            for p in cells[ci].paragraphs:
                                for r in p.runs:
                                    r.bold = True
            continue
        # Headings
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            doc.add_heading(_strip_inline(m.group(2)), level=min(len(m.group(1)), 4))
            i += 1
            continue
        # Code fence
        if s.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            continue
        # Bullets / numbered
        if re.match(r"^[-*+]\s+", s):
            doc.add_paragraph(_strip_inline(s[2:]), style="List Bullet")
            i += 1
            continue
        if re.match(r"^\d+[.)]\s+", s):
            doc.add_paragraph(_strip_inline(re.sub(r"^\d+[.)]\s+", "", s)), style="List Number")
            i += 1
            continue
        # Paragraph — preserve **bold** as runs.
        p = doc.add_paragraph()
        for j, part in enumerate(_BOLD_RE.split(s)):
            run = p.add_run(part.replace("`", ""))
            if j % 2 == 1:
                run.bold = True
        i += 1
    doc.save(path)


def _render_xlsx(md: str, title: str, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    tables = _parse_tables(md)
    if tables:
        for n, table in enumerate(tables, 1):
            ws = wb.create_sheet(title=(_slug(title, "sheet")[:25] if len(tables) == 1 else f"Table{n}"))
            for row in table:
                ws.append(row)
            for c in ws[1]:
                c.font = Font(bold=True)
            # Auto-ish column width.
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    else:
        ws = wb.create_sheet(title=_slug(title, "sheet")[:25])
        for line in md.splitlines():
            if line.strip():
                ws.append([_strip_inline(line)])
    wb.save(path)


def _render_pptx(md: str, title: str, path: str) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    # Title slide.
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = title or "Presentation"

    cur = None  # current body text-frame
    for line in md.splitlines():
        s = line.strip()
        if not s:
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = _strip_inline(m.group(2))
            cur = slide.placeholders[1].text_frame
            cur.clear()
            cur._first = True
            continue
        text = _strip_inline(re.sub(r"^[-*+]\s+|^\d+[.)]\s+", "", s))
        if cur is None:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title or "Details"
            cur = slide.placeholders[1].text_frame
            cur.clear()
            cur._first = True
        if getattr(cur, "_first", False):
            cur.paragraphs[0].text = text
            cur._first = False
        else:
            p = cur.add_paragraph()
            p.text = text
            p.level = 1 if re.match(r"^\s+", line) else 0
    prs.save(path)


_RENDERERS = {"docx": _render_docx, "xlsx": _render_xlsx, "pptx": _render_pptx}


# ── Input parsing ──────────────────────────────────────────────────────────
def _parse_input(content: str) -> Tuple[Optional[str], str, str, Optional[str]]:
    """Return (format, title, markdown, doc_id) from JSON or line-based input."""
    raw = (content or "").strip()
    fmt = title = markdown = doc_id = None
    # JSON form (native tool-calls / careful models).
    if raw.startswith("{"):
        try:
            d = json.loads(raw)
            if isinstance(d, dict):
                fmt = d.get("format") or d.get("type")
                title = d.get("title")
                markdown = d.get("content") or d.get("markdown") or d.get("body")
                doc_id = d.get("doc_id") or d.get("document_id")
        except (ValueError, TypeError):
            pass
    # Line-based form: `format:` / `title:` headers, rest is markdown.
    if fmt is None and markdown is None:
        lines = raw.splitlines()
        body_start = 0
        for idx, ln in enumerate(lines[:4]):
            low = ln.strip().lower()
            if low.startswith("format:"):
                fmt = ln.split(":", 1)[1].strip()
                body_start = idx + 1
            elif low.startswith("title:"):
                title = ln.split(":", 1)[1].strip()
                body_start = idx + 1
            else:
                break
        markdown = "\n".join(lines[body_start:]).strip()
    return (_normalize_format(fmt) if fmt else None,
            (title or "").strip(), (markdown or "").strip(), doc_id)


class ExportDocumentTool:
    async def execute(self, content: str, ctx: dict) -> Dict:
        from src.constants import DATA_DIR

        fmt, title, markdown, doc_id = _parse_input(content)
        if not fmt:
            return {"error": (
                "export_document needs a format. Use:\n"
                "```export_document\nformat: docx|xlsx|pptx\ntitle: My Title\n"
                "<markdown content>\n```\n"
                "(docx=Word, xlsx=Excel, pptx=PowerPoint). For plain markdown, "
                "use create_document instead — only export when the user asks for "
                "a specific Office format.")}

        # No inline content → export an existing document by id / the active one.
        if not markdown:
            try:
                from src.database import SessionLocal, Document
                from src.agent_tools.document_tools import _active_document_id
                db = SessionLocal()
                try:
                    target = doc_id or _active_document_id
                    q = db.query(Document)
                    doc = (q.filter(Document.id == target).first() if target
                           else q.filter(Document.owner == ctx.get("owner"))
                                 .order_by(Document.updated_at.desc()).first())
                    if doc:
                        markdown = doc.current_content or ""
                        title = title or doc.title or "export"
                finally:
                    db.close()
            except Exception as e:
                logger.debug("export: doc lookup failed: %s", e)
        if not markdown:
            return {"error": "export_document: no content to export (pass markdown content or open a document first)."}

        title = title or "export"
        try:
            out_dir = os.path.join(DATA_DIR, "exports")
            os.makedirs(out_dir, exist_ok=True)
            path = os.path.join(out_dir, _slug(title) + _EXT[fmt])
            try:
                _RENDERERS[fmt](markdown, title, path)
            except ImportError:
                return {"error": (
                    f"The library for {fmt} export isn't installed. Install with: "
                    f"pip install {'python-docx' if fmt=='docx' else 'python-pptx' if fmt=='pptx' else 'openpyxl'}")}
            size = os.path.getsize(path)
            return {
                "output": f"Exported {fmt.upper()} → {path} ({size} bytes). Open it from that path.",
                "path": path,
                "format": fmt,
                "exit_code": 0,
            }
        except Exception as e:
            logger.warning("export_document failed: %s", e)
            return {"error": f"export failed: {e}"}
